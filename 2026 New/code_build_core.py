from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd


INPUT_PATH = Path("data_panel_2018-2024.parquet")
OUTPUT_PARQUET_PATH = Path("data_core_2018-2024.parquet")
OUTPUT_XLSX_PATH = Path("data_core_2018-2024.xlsx")

PRICE_INDEX_BY_YEAR = {
    # Statistics Poland annual CPI for 2019 versus 2018 was 1.023;
    # invert it to preserve the existing 2019 = 1.0 base.
    2018: 0.977517106549,
    2019: 1.000000000,
    2020: 1.034000000,
    2021: 1.086734000,
    2022: 1.243223696,
    2023: 1.384951196,
    2024: 1.434809439,
}

EXPECTED_YEARS = {2018, 2019, 2020, 2021, 2022, 2023, 2024}
P1_LAG_SOURCE_YEAR = 2018
MAIN_ANALYSIS_START_YEAR = 2019

SECTOR_EN_MAP = {
    "budownictwo": "construction",
    "chemia": "chemicals",
    "energetyka": "energy",
    "górnictwo i hutnictwo": "mining and metallurgy",
    "handel detaliczny": "retail trade",
    "handel hurtowy": "wholesale trade",
    "media, telekomunkacja, it": "media, telecommunications, and IT",
    "motoryzacja": "automotive",
    "ochrona zdrowia i farmacja": "health and pharma",
    "paliwa": "fuels",
    "produkcja": "production",
    "transport": "transport",
    "usługi": "services",
    "żywność": "food",
}

KEY_COLUMNS = ["nip", "year"]

STATIC_COLUMNS = [
    "company",
    "rank_2019",
    "in_rank_2019",
    "pkd",
    "pkd_description",
    "sector",
    "sector_en",
    "manufacturing",
    "owner_type",
    "owner",
    "owner_num",
    "city",
    "regon",
    "krs",
    "legal_form",
    "sj",
]

ANNUAL_COLUMNS = [
    "sales",
    "operating_result",
    "profit_before_tax",
    "income_tax",
    "net_profit",
    "depreciation",
    "exports",
    "employment",
    "wages_total",
    "total_assets",
    "fixed_assets",
    "current_assets",
    "equity",
    "zobowiazania_i_rezerwy_na_zobowiazania",
    "zobowiazania_dlugoterminowe",
    "zobowiazania_krotkoterminow",
    "liabilities_provisions",
    "total_liabilities",
]

DERIVED_COLUMNS = [
    "price_index",
    "sales_real",
    "ln_sales",
    "ln_total_assets",
    "ln_employment",
    "profit_margin",
    "operating_margin",
    "export_ratio",
    "asset_turnover",
    "capital_ratio",
    "equity_multiplier",
    "roa",
    "roe",
    "depreciation_ratio",
    "wage_intensity",
    "sales_per_employee",
    "assets_per_employee",
    "sales_growth_yoy",
    "sales_real_growth_yoy",
    "sales_log_growth_yoy",
    "has_sales",
    "has_assets",
    "has_employment",
]

DROP_COLUMNS = [
    "rank_2020",
    "rank_2021",
    "rank_2022",
    "rank_2023",
    "rank_2024",
]

IGNORED_INPUT_COLUMNS = ["gpw", "przychody"]

CORE_COLUMNS = KEY_COLUMNS + STATIC_COLUMNS + ANNUAL_COLUMNS + DERIVED_COLUMNS
INPUT_REQUIRED_COLUMNS = [
    column
    for column in KEY_COLUMNS + STATIC_COLUMNS + ANNUAL_COLUMNS
    if column
    not in {
        "in_rank_2019",
        "owner",
        "owner_num",
        "sector_en",
        "manufacturing",
        "income_tax",
        "zobowiazania_dlugoterminowe",
        "zobowiazania_krotkoterminow",
    }
] + DROP_COLUMNS

INTEGER_COLUMNS = ["year", "rank_2019", "in_rank_2019", "regon", "krs", "pkd", "owner_type"]
FLOAT_COLUMNS = [
    "sales",
    "operating_result",
    "profit_before_tax",
    "income_tax",
    "net_profit",
    "depreciation",
    "exports",
    "employment",
    "wages_total",
    "total_assets",
    "fixed_assets",
    "current_assets",
    "equity",
    "zobowiazania_i_rezerwy_na_zobowiazania",
    "zobowiazania_dlugoterminowe",
    "zobowiazania_krotkoterminow",
    "liabilities_provisions",
    "total_liabilities",
]

STRING_COLUMNS = [
    "nip",
    "company",
    "city",
    "legal_form",
    "pkd_description",
    "sector",
    "sector_en",
    "sj",
]

UNAVAILABLE_SOURCE_COLUMNS = [
    "business_start_year",
    "incorporation_year_krs",
    "income_tax",
    "zobowiazania_dlugoterminowe",
    "zobowiazania_krotkoterminow",
]

STABLE_DESCRIPTOR_SOURCE_COLUMNS = [
    "company",
    "rank_2019",
    "pkd",
    "pkd_description",
    "owner_type",
    "city",
    "regon",
    "krs",
    "legal_form",
    "sj",
]

def load_input_data(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")
    return pd.read_parquet(path)


def first_non_missing(series: pd.Series):
    non_missing = series.dropna()
    if non_missing.empty:
        return pd.NA
    return non_missing.iloc[0]


def fill_stable_descriptors(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    source_rows = df.loc[df["year"].between(MAIN_ANALYSIS_START_YEAR, 2024)]
    for column in STABLE_DESCRIPTOR_SOURCE_COLUMNS:
        if column not in df.columns:
            continue
        values_by_nip = source_rows.groupby("nip", dropna=False)[column].agg(first_non_missing)
        df[column] = df[column].fillna(df["nip"].map(values_by_nip))

    return df


def validate_expected_columns(df: pd.DataFrame) -> None:
    expected = set(INPUT_REQUIRED_COLUMNS)
    missing = sorted(expected.difference(df.columns))
    if missing:
        raise ValueError(f"Input file is missing expected columns: {missing}")


def add_unavailable_canonical_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for column in UNAVAILABLE_SOURCE_COLUMNS:
        if column in CORE_COLUMNS and column not in df.columns:
            df[column] = pd.NA
    return df


def fill_sales_from_source_revenue(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    df = df.copy()
    fill_mask = df["sales"].isna() & df["przychody"].notna()
    df.loc[fill_mask, "sales"] = df.loc[fill_mask, "przychody"]
    return df, int(fill_mask.sum())


def normalize_strings(df: pd.DataFrame) -> pd.DataFrame:
    for column in STRING_COLUMNS:
        series = df[column].astype("string").str.strip()
        df[column] = series.replace({"": pd.NA, "None": pd.NA, "nan": pd.NA})
    return df


def coerce_numeric_types(df: pd.DataFrame) -> pd.DataFrame:
    for column in INTEGER_COLUMNS:
        df[column] = pd.to_numeric(df[column], errors="coerce").round().astype("Int64")

    for column in FLOAT_COLUMNS:
        df[column] = pd.to_numeric(df[column], errors="coerce").astype(float)

    return df


def add_rank_indicators(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["rank_2019"] = pd.to_numeric(df["rank_2019"], errors="coerce").round().astype("Int64")
    df["in_rank_2019"] = df["rank_2019"].notna().astype("Int64")
    return df


def safe_log(series: pd.Series) -> pd.Series:
    positive = series.where(series > 0)
    return np.log(positive)


def safe_ratio(numerator: pd.Series, denominator: pd.Series) -> pd.Series:
    denominator = denominator.where(denominator != 0)
    return numerator.divide(denominator)


def create_sector_en(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    df = df.copy()
    sector_clean = df["sector"].astype("string").str.strip()
    sector_key = sector_clean.str.lower()
    df["sector"] = sector_clean
    df["sector_en"] = sector_key.map(SECTOR_EN_MAP).astype("string")
    unmatched = sorted(
        sector_clean.loc[sector_clean.notna() & df["sector_en"].isna()]
        .dropna()
        .unique()
        .tolist()
    )
    return df, unmatched


def build_derived_variables(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    df = df.copy()
    df, unmatched_sector_values = create_sector_en(df)

    owner_type_string = (
        pd.to_numeric(df["owner_type"], errors="coerce")
        .round()
        .astype("Int64")
        .astype("string")
    )
    df["owner"] = np.where(owner_type_string.str.startswith("5", na=False), "Foreign", "Domestic")
    df["owner"] = pd.Series(df["owner"], index=df.index, dtype="string")
    df["owner_num"] = np.where(df["owner"] == "Foreign", 1, 0)
    df["owner_num"] = pd.Series(df["owner_num"], index=df.index, dtype="Int64")

    pkd_prefix = (
        pd.to_numeric(df["pkd"], errors="coerce")
        .round()
        .astype("Int64")
        .astype("string")
        .str.zfill(4)
        .str[:2]
    )
    pkd_section = pd.to_numeric(pkd_prefix, errors="coerce").astype("Int64")
    df["manufacturing"] = pkd_section.between(10, 33, inclusive="both").astype("Int64")

    df["price_index"] = df["year"].map(PRICE_INDEX_BY_YEAR).astype(float)
    df["sales_real"] = safe_ratio(df["sales"], df["price_index"])

    df["ln_sales"] = safe_log(df["sales"])
    df["ln_total_assets"] = safe_log(df["total_assets"])
    df["ln_employment"] = safe_log(df["employment"])

    df["profit_margin"] = safe_ratio(df["net_profit"], df["sales"])
    df["operating_margin"] = safe_ratio(df["operating_result"], df["sales"])
    df["export_ratio"] = safe_ratio(df["exports"], df["sales"])
    df["asset_turnover"] = safe_ratio(df["sales"], df["total_assets"])
    df["capital_ratio"] = safe_ratio(df["equity"], df["total_assets"])
    df["equity_multiplier"] = safe_ratio(df["total_assets"], df["equity"])
    df["roa"] = safe_ratio(df["net_profit"], df["total_assets"])
    df["roe"] = safe_ratio(df["net_profit"], df["equity"])
    df["depreciation_ratio"] = safe_ratio(df["depreciation"], df["total_assets"])
    df["wage_intensity"] = safe_ratio(df["wages_total"], df["sales"])

    df["sales_per_employee"] = safe_ratio(df["sales"], df["employment"])
    df["assets_per_employee"] = safe_ratio(df["total_assets"], df["employment"])

    df["has_sales"] = df["sales"].gt(0).fillna(False).astype("Int64")
    df["has_assets"] = df["total_assets"].gt(0).fillna(False).astype("Int64")
    df["has_employment"] = df["employment"].gt(0).fillna(False).astype("Int64")

    df = df.sort_values(KEY_COLUMNS, kind="mergesort")
    lag_sales = df.groupby("nip", sort=False)["sales"].shift(1)
    lag_sales_real = df.groupby("nip", sort=False)["sales_real"].shift(1)
    lag_ln_sales = df.groupby("nip", sort=False)["ln_sales"].shift(1)

    df["sales_growth_yoy"] = np.where(lag_sales > 0, df["sales"] / lag_sales - 1, np.nan)
    df["sales_real_growth_yoy"] = np.where(lag_sales_real > 0, df["sales_real"] / lag_sales_real - 1, np.nan)
    df["sales_log_growth_yoy"] = df["ln_sales"] - lag_ln_sales

    return df, unmatched_sector_values


def select_core_columns(df: pd.DataFrame) -> pd.DataFrame:
    extra_columns = sorted(
        set(df.columns).difference(CORE_COLUMNS + DROP_COLUMNS + IGNORED_INPUT_COLUMNS)
    )
    if extra_columns:
        raise ValueError(f"Unexpected columns found in input: {extra_columns}")
    return df.loc[:, CORE_COLUMNS].copy()


def drop_invalid_keys(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(subset=KEY_COLUMNS).copy()
    df["year"] = df["year"].astype("Int64")
    return df


def deduplicate_firm_year(df: pd.DataFrame) -> pd.DataFrame:
    if not df.duplicated(subset=KEY_COLUMNS).any():
        return df

    working = df.copy()
    value_columns = [column for column in working.columns if column not in KEY_COLUMNS]
    working["_non_null_count"] = working[value_columns].notna().sum(axis=1)

    working = working.sort_values(
        by=KEY_COLUMNS + ["_non_null_count", "company"],
        ascending=[True, True, False, True],
        kind="mergesort",
    )

    working = working.drop_duplicates(subset=KEY_COLUMNS, keep="first")
    working = working.drop(columns="_non_null_count")
    return working


def sort_panel(df: pd.DataFrame) -> pd.DataFrame:
    return df.sort_values(KEY_COLUMNS, kind="mergesort").reset_index(drop=True)


def validate_output(df: pd.DataFrame) -> None:
    if df.duplicated(subset=KEY_COLUMNS).any():
        raise ValueError("Output still contains duplicate (nip, year) rows.")

    if list(df.columns) != CORE_COLUMNS:
        raise ValueError("Output columns do not match the expected canonical schema.")

    output_years = set(df["year"].dropna().astype(int).unique())
    if output_years != EXPECTED_YEARS:
        raise ValueError(f"Output years do not match the expected 2018-2024 range: {sorted(output_years)}")

    rows_2018 = df["year"].eq(P1_LAG_SOURCE_YEAR)
    if not rows_2018.any():
        raise ValueError("Output does not contain 2018 rows needed for P1 lag growth.")
    if df.loc[rows_2018, "sales"].notna().sum() == 0:
        raise ValueError("Output has no observed 2018 sales values for P1 lag growth.")
    if df.loc[rows_2018, "price_index"].isna().any():
        raise ValueError("Output has missing 2018 price_index values.")
    if "sector" not in df.columns or "sector_en" not in df.columns:
        raise ValueError("Output must contain sector and sector_en.")
    if df["sector_en"].dropna().eq("production").sum() == 0:
        raise ValueError("Output sector_en has no production observations.")


def write_outputs(df: pd.DataFrame, parquet_path: Path, xlsx_path: Path) -> None:
    df.to_parquet(parquet_path, index=False)
    df.to_excel(xlsx_path, index=False, engine="xlsxwriter")
    format_excel_output(xlsx_path)


def format_excel_output(xlsx_path: Path = OUTPUT_XLSX_PATH) -> None:
    """Apply Period-style analysis formatting without changing existing cells.

    Callable separately for formatting-only refreshes. A full build still takes
    its data exclusively from the canonical builder, never from Excel edits.
    The large Core workbook uses a bounded native styling pass because the
    artifact-tool full-size authoring path exceeded its memory budget previously.
    """
    from copy import deepcopy
    from math import ceil
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.views import Selection
    from openpyxl.utils import get_column_letter
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    import re
    import io

    # Styling libraries may reserialize a float with fewer significant digits.
    # Retain the original numeric XML tokens so a format-only save cannot alter
    # a number by even one floating-point unit. Strings are preserved by the
    # workbook loader; no numeric/formula calculation is performed here.
    with ZipFile(xlsx_path) as source_archive:
        namespaces = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        original_sheet = ET.fromstring(source_archive.read("xl/workbook.xml")).find("s:sheets/s:sheet", namespaces)
        relation_id = original_sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        relationships = ET.fromstring(source_archive.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in relationships if r.attrib["Id"] == relation_id)
        source_part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        original_numbers = {}
        tag_c = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"
        tag_v = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v"
        for _, element in ET.iterparse(io.BytesIO(source_archive.read(source_part)), events=("end",)):
            if element.tag == tag_c:
                value = element.find(tag_v)
                if element.get("t", "n") in {"n", "b"} and value is not None and value.text is not None:
                    original_numbers[element.attrib["r"]] = value.text
                element.clear()

    workbook = load_workbook(xlsx_path)
    worksheet = workbook.worksheets[0]
    columns = [cell.value for cell in worksheet[1]]
    if columns != CORE_COLUMNS:
        workbook.close()
        raise ValueError("Core Excel schema differs from canonical columns; formatting cancelled.")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(right=Side(style="thin", color="FFFFFF"))
    wrap_columns = {"company", "pkd_description", "sector", "sector_en", "legal_form"}
    integer_columns = set(INTEGER_COLUMNS) | {"manufacturing", "owner_num", "has_sales", "has_assets", "has_employment"}
    percentages = {"profit_margin", "operating_margin", "export_ratio", "capital_ratio", "roa", "roe", "depreciation_ratio", "wage_intensity", "sales_growth_yoy", "sales_real_growth_yoy"}
    monetary = set(ANNUAL_COLUMNS) - {"employment"}
    formats, alignments, widths = {}, {}, {}
    for index, name in enumerate(columns, 1):
        is_text = name in STRING_COLUMNS or name == "owner"
        number_format = "@" if is_text else "0" if name in integer_columns else "0.00%" if name in percentages else "#,##0.00" if name in monetary or name in {"employment", "sales_real", "sales_per_employee", "assets_per_employee"} else "0.0000"
        width = 15 if name == "nip" else 44 if name == "company" else 52 if name == "pkd_description" else 28 if name in wrap_columns else min(max(len(name) + 2, 16), 26) if not is_text else min(max(len(name) + 2, 18), 35)
        worksheet.column_dimensions[get_column_letter(index)].width = width
        formats[index], widths[name] = number_format, width
        alignments[index] = Alignment(horizontal="left" if is_text else "right", vertical="center", wrap_text=name in wrap_columns)
        cell = worksheet.cell(1, index)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, header_alignment, header_border
    worksheet.row_dimensions[1].height = 42
    for row in worksheet.iter_rows(min_row=2):
        lines = 1
        for index, cell in enumerate(row, 1):
            cell.font, cell.alignment, cell.number_format = body_font, alignments[index], formats[index]
            name = columns[index - 1]
            if name in wrap_columns and cell.value is not None:
                lines = max(lines, sum(max(1, ceil(len(part) / (widths[name] - 3))) for part in str(cell.value).split("\n")))
        worksheet.row_dimensions[row[0].row].height = 15 * lines
    worksheet.sheet_properties.tabColor = "5B9BD5"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90
    # Core starts with nip/year/company, rather than Period's nip/company.
    worksheet.freeze_panes = "D2"
    worksheet.sheet_view.topLeftCell = "A1"
    worksheet.sheet_view.selection = [
        Selection(pane="topRight", activeCell="D1", sqref="D1"),
        Selection(pane="bottomLeft", activeCell="A2", sqref="A2"),
        Selection(pane="bottomRight", activeCell="D2", sqref="D2"),
    ]
    reference = f"A1:{get_column_letter(len(columns))}{worksheet.max_row}"
    existing = worksheet.tables.get("CoreData")
    if existing is None:
        if worksheet.tables:
            workbook.close()
            raise ValueError("Unexpected existing Core Excel table; formatting cancelled to preserve its structure.")
        table = Table(displayName="CoreData", ref=reference)
        table.tableColumns = [TableColumn(id=i, name=name) for i, name in enumerate(columns, 1)]
        worksheet.add_table(table)
    else:
        table = existing
        table.ref = reference
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    # Transfer any existing worksheet filter conditions into the native table.
    prior_filter = worksheet.auto_filter if worksheet.auto_filter.ref else table.autoFilter
    table.autoFilter = AutoFilter(ref=reference,
                                 filterColumn=deepcopy(prior_filter.filterColumn) if prior_filter is not None else [],
                                 sortState=deepcopy(prior_filter.sortState) if prior_filter is not None else None)
    worksheet.auto_filter.ref = None
    temporary = xlsx_path.with_name(xlsx_path.stem + ".formatting.xlsx")
    workbook.save(temporary)
    workbook.close()
    precise_temporary = xlsx_path.with_name(xlsx_path.stem + ".exact-formatting.xlsx")
    with ZipFile(temporary) as formatted_archive, ZipFile(precise_temporary, "w") as precise_archive:
        for entry in formatted_archive.infolist():
            content = formatted_archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                xml = content.decode("utf-8")
                def preserve_number(match):
                    cell = match.group(0)
                    coordinate = re.search(r'\br="([^"]+)"', cell)
                    raw = original_numbers.get(coordinate.group(1)) if coordinate else None
                    return re.sub(r"<v>[^<]*</v>", lambda _: "<v>" + raw + "</v>", cell, count=1) if raw is not None else cell
                content = re.sub(r"<c\b[^>]*>.*?</c>", preserve_number, xml, flags=re.DOTALL).encode("utf-8")
            precise_archive.writestr(entry, content)
    precise_temporary.replace(xlsx_path)
    temporary.unlink()
    print(f"Core Excel formatting complete: {worksheet.max_row - 1:,} rows, {len(columns)} columns; CoreData table; identifiers/header frozen; values retained.")


def build_core_panel(input_path: Path = INPUT_PATH) -> tuple[pd.DataFrame, dict]:
    df = load_input_data(input_path)
    input_row_count = len(df)
    validate_expected_columns(df)
    df = add_unavailable_canonical_columns(df)
    df, sales_filled_from_przychody = fill_sales_from_source_revenue(df)
    df = fill_stable_descriptors(df)
    df = add_rank_indicators(df)
    df = coerce_numeric_types(df)
    df, unmatched_sector_values = build_derived_variables(df)
    df = select_core_columns(df)
    df = normalize_strings(df)
    df = drop_invalid_keys(df)
    df = deduplicate_firm_year(df)
    df = sort_panel(df)
    validate_output(df)
    metadata = {
        "input_row_count": input_row_count,
        "output_row_count": len(df),
        "unique_firms": df["nip"].nunique(),
        "unavailable_source_columns": UNAVAILABLE_SOURCE_COLUMNS,
        "sales_filled_from_przychody": sales_filled_from_przychody,
        "firms_without_sector": int(
            df.groupby("nip")["sector"].apply(lambda series: series.notna().any()).eq(False).sum()
        ),
        "sector_missing_rate": float(df["sector"].isna().mean()),
        "sector_en_missing_rate": float(df["sector_en"].isna().mean()),
        "sector_en_values": sorted(df["sector_en"].dropna().astype(str).unique().tolist()),
        "unmatched_sector_values": unmatched_sector_values,
        "duplicate_firm_year_count": int(df.duplicated(KEY_COLUMNS).sum()),
        "output_years": sorted(df["year"].dropna().astype(int).unique().tolist()),
        "sales_2018_present_count": int(
            df.loc[df["year"].eq(P1_LAG_SOURCE_YEAR), "sales"].notna().sum()
        ),
        "rows_2018": int(df["year"].eq(P1_LAG_SOURCE_YEAR).sum()),
        "nonmissing_2018_annual_columns": {
            column: int(df.loc[df["year"].eq(P1_LAG_SOURCE_YEAR), column].notna().sum())
            for column in ANNUAL_COLUMNS
        },
    }
    return df, metadata


def print_build_summary(df: pd.DataFrame, metadata: dict) -> None:
    print("Core panel build complete")
    print(f"Input row count: {metadata['input_row_count']:,}")
    print(f"Output row count: {metadata['output_row_count']:,}")
    print(f"Number of unique firms: {metadata['unique_firms']:,}")
    print(f"Duplicate (nip, year) rows: {metadata['duplicate_firm_year_count']:,}")
    print(f"Output years: {metadata['output_years']}")
    print(
        "2018 sales is present: "
        f"{metadata['sales_2018_present_count']:,} of {metadata['rows_2018']:,} rows"
    )
    print(
        "2018 is retained only as the sales base for 2019 P1 lag growth; "
        "main growth, trajectories, SGrowth_NR, and start covariates begin in 2019."
    )
    print(
        "Non-missing 2018 annual-column counts "
        f"(missing non-sales annual values are allowed): {metadata['nonmissing_2018_annual_columns']}"
    )
    print(
        "Columns omitted because they are unavailable in the current source: "
        f"{metadata['unavailable_source_columns']}"
    )
    print(
        "Missing sales values filled from source przychody: "
        f"{metadata['sales_filled_from_przychody']:,}"
    )
    print("sector is taken directly from data_panel_2018-2024.parquet.")
    print(f"Firms without an available sector descriptor: {metadata['firms_without_sector']:,}")
    print(f"sector missing rate: {metadata['sector_missing_rate']:.6f}")
    print(f"sector_en missing rate: {metadata['sector_en_missing_rate']:.6f}")
    print(f"sector_en values: {metadata['sector_en_values']}")
    print(f"Unmatched sector values: {metadata['unmatched_sector_values']}")
    print(f"sector_en present and used as categorical control: {'sector_en' in df.columns}")
    print("Final columns:")
    for column in df.columns:
        print(column)
    print(f"Final column count: {len(df.columns)}")
    print("Missing counts for key variables:")
    for column in ["sales", "sales_real", "price_index"]:
        print(f"{column}: {int(df[column].isna().sum()):,}")
    print("Summary statistics for key outcomes:")
    print(
        df[["sales", "sales_real", "sales_growth_yoy", "sales_real_growth_yoy"]]
        .describe()
        .to_string()
    )


if __name__ == "__main__":
    core_panel, build_metadata = build_core_panel()
    write_outputs(core_panel, OUTPUT_PARQUET_PATH, OUTPUT_XLSX_PATH)
    print_build_summary(core_panel, build_metadata)
