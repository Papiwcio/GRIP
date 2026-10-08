"""Meaningful synthetic and actual-data acceptance tests for the independent audit."""
from __future__ import annotations

import csv
import hashlib
from pathlib import Path
import tempfile
import unittest

import numpy as np
import pandas as pd

import code_audit_data_quality as audit
import code_build_core as core_builder
import code_build_period as period_builder


def fixtures(changes=None):
    """Produce disposable data with existing producer functions, never canonical files.

    Tests exercise the independent auditor against actual producer conventions,
    then deliberately alter source values, gaps or stored calculations.
    """
    rows = []
    for year in range(2018, 2025):
        row = dict.fromkeys(core_builder.CORE_COLUMNS, np.nan)
        row.update(nip="1234567890", year=year, company="Synthetic test firm", rank_2019=1,
                   in_rank_2019=1, pkd=1010, owner_type=1, sector="produkcja", city="Test",
                   legal_form="sp. z o.o.", regon=123456789, krs=12345, sj="J")
        row.update(sales=100. * (1.1 ** (year - 2019)))
        if year >= 2019:
            row.update(operating_result=10., profit_before_tax=8., net_profit=6.,
                       depreciation=5., exports=20., employment=2.5, wages_total=15.,
                       total_assets=1000., fixed_assets=600., current_assets=400., equity=400.)
            row[{2019: "total_liabilities", 2024: "zobowiazania_i_rezerwy_na_zobowiazania"}.get(year, "liabilities_provisions")] = 600.
        if year == 2023:
            row["operating_result"] = np.nan
        if changes and year in changes:
            row.update(changes[year])
        rows.append(row)
    frame, _ = core_builder.build_derived_variables(pd.DataFrame(rows))
    frame = frame[core_builder.CORE_COLUMNS].reset_index(drop=True)
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "fixture.parquet"
        frame.to_parquet(source, index=False)
        period, *_ = period_builder.build_period_dataset(source)
    return frame, period


class AuditTests(unittest.TestCase):
    def run_frames(self, changes=None, statistical=False):
        core, period = fixtures(changes)
        c0, p0 = core.copy(deep=True), period.copy(deep=True)
        findings, result = audit.audit_frames(core, period, statistical=statistical)
        pd.testing.assert_frame_equal(core, c0)
        pd.testing.assert_frame_equal(period, p0)
        return core, period, findings, result

    def test_clean_formulas_and_fractional_fte(self):
        _, _, f, _ = self.run_frames()
        self.assertFalse(f.classification.eq("INVALID").any(), f.loc[f.classification.eq("INVALID")].to_string())
        self.assertFalse(f.rule_id.eq("A01").any())

    def test_negative_profit_and_equity_are_valid_signed_data(self):
        _, _, f, _ = self.run_frames({2020: {"net_profit": -25., "equity": -50.}})
        self.assertFalse((f.classification.eq("INVALID") & f.variable.isin(["net_profit", "equity", "profit_margin", "roe", "capital_ratio"])).any())

    def test_negative_exports_are_suspicious(self):
        _, _, f, _ = self.run_frames({2020: {"exports": -2324.}})
        row = f.loc[f.rule_id.eq("E02")].iloc[0]
        self.assertEqual(row.classification, "SUSPICIOUS")
        self.assertEqual(row.exports, -2324.)

    def test_above_100_percent_not_invalid(self):
        _, _, f, _ = self.run_frames({2019: {"exports": 204.3}})
        row = f.loc[f.rule_id.eq("E01")].iloc[0]
        self.assertEqual(row.classification, "SUSPICIOUS")
        self.assertAlmostEqual(row.calculated_value, 2.043)
        self.assertIn("export_ratio_start_P1", row.affected_variables)
        self.assertIn("FULL", row.affected_variables)

    def test_zero_sales_is_not_missing_or_invalid_raw_value(self):
        _, _, f, _ = self.run_frames({2020: {"sales": 0., "exports": 10.}})
        self.assertTrue((f.rule_id.eq("A04") & f.variable.eq("sales")).any())
        self.assertFalse((f.rule_id.eq("D05") & f.variable.eq("sales")).any())
        self.assertFalse(f.classification.eq("INVALID").any())

    def test_negative_sales_log_domain_handled(self):
        _, _, f, _ = self.run_frames({2020: {"sales": -10.}})
        self.assertTrue((f.rule_id.eq("A03") & f.variable.eq("sales")).any())
        self.assertFalse((f.rule_id.eq("A06") & f.variable.eq("ln_sales")).any())

    def test_zero_equity_does_not_produce_a_valid_roe(self):
        c, p = fixtures({2020: {"equity": 0.}})
        self.assertTrue(pd.isna(c.loc[c.year.eq(2020), "roe"]).all())
        c.loc[c.year.eq(2020), "roe"] = 7.
        f, _ = audit.audit_frames(c, p, statistical=False)
        self.assertTrue((f.rule_id.eq("A05") & f.variable.eq("roe")).any())

    def test_negative_fte_invalid_but_fractional_not(self):
        _, _, f, _ = self.run_frames({2020: {"employment": -2.5}})
        row = f.loc[f.rule_id.eq("A01")].iloc[0]
        self.assertEqual(row.classification, "INVALID")

    def test_expected_missing_2018_and_2023_blocks(self):
        _, _, f, result = self.run_frames()
        missing = f.loc[f.rule_id.eq("D05")]
        self.assertFalse(missing.year.eq(2018).any())
        self.assertFalse((missing.year.eq(2023) & missing.variable.eq("operating_result")).any())
        coverage = pd.DataFrame(result.coverage)
        self.assertTrue(coverage.loc[coverage.year.eq(2018) & coverage.variable.eq("exports"), "status"].eq("EXPECTED_UNAVAILABLE").all())

    def test_missing_sales_in_2018_and_lag_integrity(self):
        c, p = fixtures()
        c.loc[c.year.eq(2018), ["sales", "sales_real", "ln_sales"]] = np.nan
        f, _ = audit.audit_frames(c, p, statistical=False)
        self.assertTrue((f.rule_id.eq("D05") & f.year.eq(2018) & f.variable.eq("sales")).any())
        self.assertTrue((f.rule_id.eq("P02") & f.variable.eq("lag_ngrowth_log_ann_P1")).any())

    def test_2018_mapping_and_lag_formula(self):
        c, p = fixtures()
        c2018, c2019 = c.loc[c.year.eq(2018)].iloc[0], c.loc[c.year.eq(2019)].iloc[0]
        self.assertAlmostEqual(p.lag_ngrowth_log_ann_P1.iloc[0], np.log(c2019.sales / c2018.sales))
        panel = c.copy()
        panel["przychody"] = np.nan
        panel.loc[panel.year.eq(2018), "przychody"] = panel.loc[panel.year.eq(2018), "sales"]
        panel.loc[panel.year.eq(2018), "sales"] = np.nan
        f, _ = audit.audit_frames(c, p, panel, statistical=False)
        self.assertFalse((f.rule_id.eq("D11") & f.variable.eq("sales")).any())

    def test_negative_2018_sales_still_audited(self):
        _, _, f, _ = self.run_frames({2018: {"sales": -10.}})
        self.assertTrue((f.rule_id.eq("A03") & f.variable.eq("sales") & f.year.eq(2018)).any())
        self.assertFalse((f.rule_id.eq("D05") & f.year.eq(2018) & f.variable.ne("sales")).any())

    def test_negative_ratio_with_nonnegative_exports(self):
        _, _, f, _ = self.run_frames({2020: {"sales": -10., "exports": 10.}})
        self.assertTrue((f.rule_id.eq("E02") & f.variable.eq("export_ratio")).any())

    def test_duplicate_keys_all_candidates_reported(self):
        c, p = fixtures()
        c = pd.concat([c, c.iloc[[1]]], ignore_index=True)
        p = pd.concat([p, p], ignore_index=True)
        f, a = audit.audit_frames(c, p, statistical=False)
        self.assertGreaterEqual(f.rule_id.eq("D02").sum(), 4)
        self.assertTrue(any(x["status"] == "NOT_EVALUABLE" and x["rule_id"] == "P01" for x in a.contexts))

    def test_nonconsecutive_calendar_yoy_flagged(self):
        c, p = fixtures()
        c = c.loc[c.year.ne(2021)].reset_index(drop=True)
        # Simulate the current producer's previous-observed-row calculation.
        c.loc[c.year.eq(2022), "sales_growth_yoy"] = 1.1 ** 2 - 1
        f, _ = audit.audit_frames(c, p, statistical=False)
        self.assertTrue((f.rule_id.eq("T07") & f.year.eq(2022) & f.variable.eq("sales_growth_yoy")).any())

    def test_discarded_duplicate_growth_dependency_detected(self):
        c, p = fixtures()
        source = pd.concat([c, c.iloc[[1]]], ignore_index=True)
        source.loc[source.index[-1], "sales"] = 999.
        c.loc[c.year.eq(2020), "sales_growth_yoy"] = c.loc[c.year.eq(2020), "sales"].iloc[0] / 999. - 1
        f, _ = audit.audit_frames(c, p, source, statistical=False)
        self.assertTrue(f.rule_id.eq("D10").any())
        self.assertTrue((f.rule_id.eq("T07") & f.year.eq(2020)).any())

    def test_calculation_mutations_are_detected(self):
        c, p = fixtures()
        c.loc[c.year.eq(2019), "export_ratio"] += .01
        p.loc[0, "ngrowth_log_ann_P2"] += .01
        p.loc[0, "lag_ngrowth_log_ann_P3"] += .01
        p.loc[0, "nindex_2024"] += 1
        p.loc[0, "export_ratio_start_P3"] += .01
        f, _ = audit.audit_frames(c, p, statistical=False)
        for rule, variable in [("A05", "export_ratio"), ("P01", "ngrowth_log_ann_P2"), ("P02", "lag_ngrowth_log_ann_P3"), ("P03", "nindex_2024"), ("P02", "export_ratio_start_P3")]:
            self.assertTrue((f.rule_id.eq(rule) & f.variable.eq(variable)).any(), variable)

    def test_missing_field_and_infinite_numeric_value(self):
        c, p = fixtures()
        c = c.drop(columns="roa")
        c.loc[c.year.eq(2020), "exports"] = np.inf
        f, _ = audit.audit_frames(c, p, statistical=False)
        self.assertTrue((f.rule_id.eq("D01") & f.variable.eq("roa")).any())
        self.assertTrue((f.rule_id.eq("D04") & f.variable.eq("exports")).any())

    def test_reference_minimum_and_zero_dispersion(self):
        _, _, f, a = self.run_frames(statistical=True)
        self.assertFalse(f.rule_id.eq("T05").any())
        self.assertTrue(any(x["rule_id"] == "T05" and x["status"] == "NOT_EVALUABLE" for x in a.contexts))
        frame = pd.DataFrame({"nip": [str(x) for x in range(31)], "company": "Fixture", "year": 2020, "sector_en": "production", "profit_margin": [1.] * 30 + [1000.]})
        result = audit.Audit()
        audit.screen(result, frame, ["profit_margin"], "Core")
        self.assertFalse(result.rows)  # MAD=IQR=0 is not assigned an invented scale.

    def test_mad_and_iqr_fallback(self):
        frame = pd.DataFrame({"nip": [str(x) for x in range(40)], "company": "Fixture", "year": 2020, "sector_en": "production", "profit_margin": list(np.linspace(-1, 1, 39)) + [1000.]})
        result = audit.Audit()
        audit.screen(result, frame, ["profit_margin"], "Core")
        self.assertTrue(any(x["classification"] == "EXTREME" for x in result.rows))
        frame.profit_margin = [0.] * 25 + [1.] * 14 + [1000.]
        result = audit.Audit()
        audit.screen(result, frame, ["profit_margin"], "Core")
        self.assertTrue(any('IQR fallback' in x["related_values"] for x in result.rows))

    def test_stable_ids_multiple_findings_and_decision_persistence(self):
        _, _, f, _ = self.run_frames({2020: {"exports": -2324., "employment": -1.}})
        self.assertGreater(len(f), f.nip.nunique())
        self.assertFalse(f.finding_id.duplicated().any())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "decisions.csv"
            audit.merge_decisions(f, path)
            with path.open() as handle:
                rows = list(csv.DictReader(handle))
            rows[0]["decision"] = "Retain"
            rows[0]["reviewer"] = "Researcher"
            with path.open('w', newline='') as handle:
                writer = csv.DictWriter(handle, fieldnames=audit.DECISION_COLUMNS)
                writer.writeheader(); writer.writerows(rows)
            before = path.read_bytes()
            merged = audit.merge_decisions(f, path)
            self.assertEqual(before, path.read_bytes())
            self.assertEqual(merged.loc[merged.finding_id.eq(rows[0]["finding_id"]), "review_status"].iloc[0], "Retain")
            changed = f.sample(frac=1, random_state=2)
            self.assertEqual(set(f.finding_id), set(changed.finding_id))

    def test_known_actual_export_cases_and_no_manual_filter(self):
        c = pd.read_parquet(audit.ROOT / "data_core_2018-2024.parquet")
        p = pd.read_parquet(audit.ROOT / "data_period_2018-2024.parquet")
        protected = [audit.ROOT / name for name in ["data_core_2018-2024.parquet", "data_period_2018-2024.parquet", "code_config.py", "code_ols_scenarios.py"]]
        before = {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in protected}
        f, _ = audit.audit_frames(c, p, statistical=False)
        for nip, year, rule in [("5260207930",2019,"E01"),("5260207930",2020,"E01"),("6790175807",2020,"E02")]:
            rows = f.loc[f.nip.eq(nip) & f.year.eq(year) & f.rule_id.eq(rule)]
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows.classification.iloc[0], "SUSPICIOUS")
        self.assertTrue(f.existing_manual_exclusion.any())
        after = {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in protected}
        self.assertEqual(before, after)

    def test_saved_workbook_complete_native_tables_and_known_values(self):
        import json
        from zipfile import ZipFile
        import xml.etree.ElementTree as ET
        from openpyxl import load_workbook
        file = audit.ROOT / "results_data_quality_audit.xlsx"
        metadata = audit.ROOT / "results_data_quality_audit_metadata.json"
        if not file.exists() or not metadata.exists():
            self.skipTest("Run the audit first to validate the saved workbook.")
        expected = json.loads(metadata.read_text())
        workbook = load_workbook(file, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["README", "Summary", "Core_Findings", "Period_Findings", "Company_Review", "Rule_Register"])
        identifiers, known = set(), set()
        for name in ["Core_Findings", "Period_Findings"]:
            rows = workbook[name].iter_rows(min_row=4, values_only=True)
            columns = next(rows)
            n = 0
            for row in rows:
                values = dict(zip(columns, row)); n += 1
                self.assertNotIn(values["finding_id"], identifiers)
                identifiers.add(values["finding_id"])
                if values["nip"] in {"5260207930", "6790175807"} and values["rule_id"] in {"E01", "E02"}:
                    known.add((values["nip"], values["year"], values["rule_id"]))
                    self.assertIsInstance(values["sales"], (int, float))
                    self.assertIsInstance(values["exports"], (int, float))
            self.assertEqual(workbook[name].max_row, n + 4)
        workbook.close()
        self.assertEqual(len(identifiers), expected["summary"]["Total findings"])
        self.assertTrue({("5260207930",2019,"E01"),("5260207930",2020,"E01"),("6790175807",2020,"E02")}.issubset(known))
        with ZipFile(file) as zipped:
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for path in [p for p in zipped.namelist() if p.startswith("xl/tables/")]:
                table = ET.fromstring(zipped.read(path))
                self.assertIsNotNone(table.find("s:autoFilter", ns))
            for number in [3,4]:
                sheet = ET.fromstring(zipped.read(f"xl/worksheets/sheet{number}.xml"))
                self.assertEqual(sheet.find(".//s:pane", ns).get("topLeftCell"), "C5")
        self.assertTrue(expected["protected_files_unchanged"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
