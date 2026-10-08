"""Reproduce both OLS specifications and verify paired workbook architecture."""
from contextlib import redirect_stdout
from copy import deepcopy
import hashlib
import io
import os
import tempfile
from pathlib import Path
import unittest

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
import code_ols_scenarios as o


class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.protected = [p for p in Path('.').rglob('*') if p.is_file() and (p.suffix in {'.xlsx','.parquet'} or p.name == 'code_config.py') and p.name not in {'results_ols_scenarios.xlsx','results_ols_interactions.xlsx'}]
        cls.hashes = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in cls.protected}
        with redirect_stdout(io.StringIO()):
            cls.primary = o.run_period_ols_scenarios({**o.CONFIG,'include_interactions':False}, write_output=False)
            cls.extended = o.run_period_ols_scenarios({**o.CONFIG,'include_interactions':True}, write_output=False)
        cls.comparison = o.build_model_comparison(cls.primary, cls.extended)
        # The public interaction file is now the compact supplement. Preserve
        # validation of the historical full writer in a disposable workbook.
        cls.full_workbook_directory = tempfile.TemporaryDirectory(prefix='grip_full_ols_check_')
        cls.extended_workbook = str(Path(cls.full_workbook_directory.name)/'full_interactions.xlsx')
        cls.primary_workbook = str(Path(cls.full_workbook_directory.name)/'primary.xlsx')
        primary_arguments={**cls.primary['workbook_arguments'],'config':{**cls.primary['config_used'],'output_file':cls.primary_workbook,'annotation_source':'results_ols_scenarios.xlsx'}}
        o.write_scenario_workbook(**primary_arguments)
        arguments = {**cls.extended['workbook_arguments'], 'config': {**cls.extended['config_used'], 'output_file': cls.extended_workbook,
                     'annotation_source': os.environ.get('GRIP_OLS_LEGACY_WORKBOOK','archive/results_ols_interactions_2026-10-08.xlsx')}}
        o.write_scenario_workbook(**arguments,model_comparison=cls.comparison)

    @classmethod
    def tearDownClass(cls):
        cls.full_workbook_directory.cleanup()

    def pairs(self):
        for scenario in o.SCENARIOS:
            left = self.primary['scenario_results'][scenario]['model_results']
            right = self.extended['scenario_results'][scenario]['model_results']
            self.assertEqual([m['model'] for m in left], [m['model'] for m in right])
            yield from zip(left,right)

    def test_all_128_models_and_paired_observations(self):
        for report in [self.primary,self.extended]:
            self.assertEqual(report['models_estimated'],64)
            self.assertEqual(report['models_skipped'],0)
        for left,right in self.pairs():
            self.assertEqual(left['estimation_sample_ids'],right['estimation_sample_ids'])
            self.assertEqual(left['sample_sha256'],right['sample_sha256'])
            np.testing.assert_array_equal(left['result'].model.endog,right['result'].model.endog)
            self.assertEqual(left['covariance_type'],right['covariance_type'])
            self.assertEqual(left['winsorisation_lower_bound'],right['winsorisation_lower_bound'])
            self.assertEqual(left['winsorisation_upper_bound'],right['winsorisation_upper_bound'])

    def test_primary_reproduces_independent_reduced_design(self):
        for left,right in self.pairs():
            expanded = right['result']
            interaction_columns = set(right['interaction_centring'])
            design = pd.DataFrame(expanded.model.exog,index=expanded.model.data.row_labels,columns=expanded.model.exog_names).drop(columns=list(interaction_columns))
            y = pd.Series(expanded.model.endog,index=design.index)
            independent = sm.OLS(y,design).fit(cov_type=left['covariance_type'])
            current = left['result']
            self.assertEqual(current.model.exog_names,list(design.columns))
            np.testing.assert_allclose(current.params,independent.params,rtol=1e-10,atol=1e-10)
            np.testing.assert_allclose(current.pvalues,independent.pvalues,rtol=1e-9,atol=1e-10)
            np.testing.assert_allclose(current.cov_params(),independent.cov_params(),rtol=1e-9,atol=1e-10)
            np.testing.assert_allclose(current.fittedvalues,independent.fittedvalues,rtol=1e-10,atol=1e-10)
            self.assertAlmostEqual(current.rsquared,independent.rsquared,places=12)

    def test_saved_models_match_both_specifications(self):
        for file,report in [('results_ols_scenarios.xlsx',self.primary),(self.extended_workbook,self.extended)]:
            saved = pd.read_excel(file,sheet_name='Coefficients_Long')
            cols=['scenario','model','raw_variable']
            expected=report['coefficients_long_df']
            merged=expected.merge(saved,on=cols,suffixes=('_expected','_saved'),validate='one_to_one')
            self.assertEqual(len(merged),len(expected))
            for field in ['coefficient','std error','p-value','CI lower','CI upper']:
                np.testing.assert_allclose(merged[field+'_expected'],merged[field+'_saved'],rtol=1e-10,atol=1e-10)
        self.assertFalse(self.primary['coefficients_long_df'].variable_type.eq('interaction').any())
        self.assertTrue(self.extended['coefficients_long_df'].variable_type.eq('interaction').any())

    def test_existing_centred_interaction_results_reproduced(self):
        path=os.environ.get('GRIP_OLS_LEGACY_WORKBOOK')
        if not path:
            self.skipTest('Set GRIP_OLS_LEGACY_WORKBOOK to the preserved pre-split workbook for migration verification.')
        old=pd.read_excel(path,sheet_name='Coefficients_Long')
        current=self.extended['coefficients_long_df']
        merged=old.merge(current,on=['scenario','model','raw_variable'],suffixes=('_old','_new'),validate='one_to_one')
        self.assertEqual(len(merged),len(old))
        for field in ['coefficient','std error','t-stat','p-value','CI lower','CI upper']:
            np.testing.assert_allclose(merged[field+'_old'],merged[field+'_new'],rtol=1e-10,atol=1e-10)
        summary_old=pd.read_excel(path,sheet_name='Model_Summary_Long')
        summary_new=self.extended['summary_long_df']
        pairs=summary_old.merge(summary_new,on=['scenario','model'],suffixes=('_old','_new'),validate='one_to_one')
        for field in ['observations','R_squared','adjusted_R_squared']:
            np.testing.assert_allclose(pairs[field+'_old'],pairs[field+'_new'],rtol=1e-10,atol=1e-10)

    def test_working_and_technical_sheets_and_comparison(self):
        # Check formatter defaults on disposable outputs. The live primary may
        # have researcher-adjusted panes and must not be reformatted by a test.
        for file,extended in [(self.primary_workbook,False),(self.extended_workbook,True)]:
            with open(file,'rb') as handle:
                workbook=load_workbook(handle,read_only=False,data_only=True)
            names=o.OUTPUT_SHEETS.copy()
            if extended:names.insert(names.index('AUDIT_AND_TECHNICAL_TABS'),'Model_Comparison')
            self.assertEqual(workbook.sheetnames,names)
            self.assertEqual(workbook['Compare_Main'].freeze_panes,'C2')
            if extended:
                self.assertEqual(workbook['Model_Comparison'].freeze_panes,'D2')
                self.assertTrue(workbook['Model_Comparison'].auto_filter.ref)
            workbook.close()
        saved=pd.read_excel(self.extended_workbook,sheet_name='Model_Comparison')
        self.assertEqual(len(saved),64)
        self.assertTrue(saved.identical_observations.all())
        np.testing.assert_allclose(saved.delta_R2,saved.R2_with-saved.R2_without,rtol=1e-10,atol=1e-12)
        self.assertTrue(saved.delta_R2.ge(-1e-12).all())
        for column in self.comparison:
            if pd.api.types.is_numeric_dtype(self.comparison[column]) and column != 'identical_observations':
                np.testing.assert_allclose(saved[column],self.comparison[column],rtol=1e-10,atol=1e-10)

    def test_report_switch_and_missing_values_do_not_leak(self):
        data=pd.read_parquet(o.CONFIG['input_file']).head(150).copy()
        data.loc[data.index[0],'export_ratio_start_P1']=np.nan
        configurations=[o.normalise_config({**o.CONFIG,'include_interactions':flag}) for flag in [False,True]]
        samples=[]
        for config in configurations:
            models=o.build_models(config)
            frame=o.add_interaction_columns(data,config,models)
            samples.append(o.get_estimation_sample(frame,config,models['P1']).index)
        pd.testing.assert_index_equal(*samples)
        self.assertTrue(o.get_active_interactions())  # Legacy imports retain active metadata.
        self.assertFalse(o.get_active_interactions(configurations[0]))
        metadata=deepcopy(o.INTERACTION_METADATA)
        try:
            o.INTERACTION_METADATA['size_owner']={'variables':['ln_sales','owner_num'],'display_name':'size × owner','interpretation':'test','standardise':True,'include':True}
            self.assertFalse(o.get_active_interactions(configurations[0]))
            self.assertEqual(len(o.get_active_interactions(configurations[1])),2)
        finally:
            o.INTERACTION_METADATA.clear();o.INTERACTION_METADATA.update(metadata)

    def test_mismatched_pair_cannot_be_published(self):
        broken=deepcopy(self.extended)
        scenario=next(iter(o.SCENARIOS))
        broken['scenario_results'][scenario]['model_results'][0]['estimation_sample_ids']=('wrong firm',)
        with self.assertRaisesRegex(ValueError,'firm IDs differ'):
            o.build_model_comparison(self.primary,broken)

    def test_robust_covariance_uses_the_same_engine(self):
        for report in [self.primary,self.extended]:
            scenario=next(iter(o.SCENARIOS))
            details=report['scenario_results'][scenario]
            config={**report['config_used'],'covariance_type':'HC3'}
            levels=o.get_categorical_levels(details['filtered_df'],config)
            spec=details['models']['P1']
            variant=o.get_model_variants(config)[0]
            result=o.run_model_variant('P1',spec,variant,details['filtered_df'],levels,details['variable_registry'],config)
            self.assertEqual(result['result'].cov_type,'HC3')
            self.assertEqual(result['estimation_sample_ids'],details['model_results'][0]['estimation_sample_ids'])

    def test_migration_annotations_preserved_in_interaction_report(self):
        path=os.environ.get('GRIP_OLS_LEGACY_WORKBOOK')
        if not path:self.skipTest('Optional pre-split migration snapshot not provided.')
        def notes(file):
            workbook=load_workbook(file,read_only=True,data_only=False)
            result={}
            for name in ['Compare_Main','Compare_Raw']:
                for row in workbook[name].iter_rows(min_row=2):
                    for index,cell in enumerate(row):
                        if index>=10 and cell.value is not None:result[name,row[0].value,row[1].value,index]=cell.value
            workbook.close()
            return result
        self.assertEqual(notes(path),notes(self.extended_workbook))

    def test_source_data_and_other_outputs_unchanged(self):
        self.assertEqual(self.hashes,{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in self.protected})


if __name__=='__main__':
    unittest.main(verbosity=2)
