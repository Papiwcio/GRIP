"""Acceptance checks for the compact separate interaction specifications."""
from contextlib import redirect_stdout
import io
import hashlib
import os
from pathlib import Path
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import code_ols_interactions as report
import code_ols_scenarios as engine


class CompactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.protected=[p for p in Path('.').rglob('*') if p.is_file() and (p.suffix in {'.xlsx','.parquet'} or p.name=='code_config.py') and p.name!='results_ols_interactions.xlsx']
        cls.hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in cls.protected}
        with redirect_stdout(io.StringIO()):
            cls.a=report.fit_specification('export_size')
            cls.b=report.fit_specification('profit_manufacturing')
        cls.summary=report.summary_table(cls.a,cls.b)

    def test_samples_variants_and_separate_terms(self):
        self.assertEqual(list(self.a['scenarios']),report.SCENARIOS)
        self.assertEqual(list(self.b['scenarios']),['ALL','RANK2019'])
        for specification,data in [('A',self.a),('B',self.b)]:
            for details in data['scenarios'].values():
                self.assertEqual(details['models_estimated'],16)  # Raw verification retained too.
                self.assertEqual(details['models_skipped'],0)
                for m in details['model_results']:
                    names=m['result'].model.exog_names
                    product=[x for x in names if '_x_' in x]
                    self.assertEqual(len(product),1)
                    self.assertIn('export_ratio_x_ln_sales' if specification=='A' else 'profit_margin_x_manufacturing',product[0])
                    if specification=='B':self.assertIn('manufacturing',names)

    def test_export_results_match_prior_full_workbook(self):
        # Historical workbook membership predates the approved Stage 2 exclusions.
        # Numerical equivalence is tested on identical IDs against the preserved engine.
        from code_check_common_samples import regression_math_equivalence
        from code_common_samples import current
        regression_math_equivalence(current())

    def test_one_profitability_scale_and_identification(self):
        for details in self.b['scenarios'].values():
            for model in details['model_results']:
                source,product=report.verify_profit_identifiability(model)
                if model['standardised_model']=='Yes':
                    x=pd.DataFrame(model['result'].model.exog,columns=model['result'].model.exog_names)
                    np.testing.assert_allclose(x[product],x[source]*x.manufacturing,atol=1e-12,rtol=1e-12)
                    self.assertEqual(set(x.manufacturing),{0.,1.})

    def test_conditional_slopes_and_full_covariance(self):
        for scenario,details in self.b['scenarios'].items():
            for model in details['model_results']:
                if model['standardised_model']!='Yes':continue
                source,product=report.verify_profit_identifiability(model)
                fit=model['result']
                row=self.summary.loc[self.summary.interaction_specification.eq('Profitability × manufacturing') & self.summary['sample'].eq(scenario) & self.summary.period.eq(model['period']) & self.summary.model_specification.eq('Standardised winsorised' if model['winsorised']=='Yes' else 'Standardised baseline')].iloc[0]
                self.assertAlmostEqual(row.profitability_non_manufacturing,float(fit.params[source]))
                self.assertAlmostEqual(row.profitability_manufacturing,float(fit.params[source]+fit.params[product]))
                self.assertAlmostEqual(row.interaction_coefficient,float(fit.params[product]))
                weights=np.zeros(len(fit.params));weights[list(fit.params.index).index(source)]=1;weights[list(fit.params.index).index(product)]=1
                contrast=fit.t_test(weights)
                self.assertAlmostEqual(row.profitability_manufacturing_p,float(np.asarray(contrast.pvalue).squeeze()),places=12)
                np.testing.assert_allclose([row.profitability_manufacturing_ci_low,row.profitability_manufacturing_ci_high],contrast.conf_int().ravel(),rtol=1e-10,atol=1e-10)

    def test_matched_reduced_models_keep_main_effects(self):
        for data in [self.a,self.b]:
            for scenario,details in data['scenarios'].items():
                for model in details['model_results']:
                    if model['standardised_model']!='Yes':continue
                    reduced=report.paired_reduced_model(data,scenario,model)
                    self.assertEqual(model['estimation_sample_ids'],reduced['estimation_sample_ids'])
                    np.testing.assert_array_equal(model['result'].model.endog,reduced['result'].model.endog)
                    if data is self.b:self.assertIn('manufacturing',reduced['result'].model.exog_names)
                    self.assertFalse(any('_x_' in x for x in reduced['result'].model.exog_names))

    def test_covariance_contrasts_follow_robust_reference(self):
        scenario='ALL';details=self.b['scenarios'][scenario];c={**self.b['config'],'covariance_type':'HC3'}
        model=engine.run_model_variant('P1',self.b['models']['P1'],engine.get_model_variants(c)[1],details['filtered_df'],engine.get_categorical_levels(details['filtered_df'],c),details['variable_registry'],c)
        source,product=report.verify_profit_identifiability(model)
        fit=model['result'];weights={source:1.,product:1.}
        calculated=report.linear_contrast(fit,weights)
        vector=[weights.get(x,0.) for x in fit.params.index]
        self.assertAlmostEqual(calculated['p'],float(np.asarray(fit.t_test(vector).pvalue).squeeze()),places=12)

    def test_four_nonempty_sheets_and_exact_familiar_tables(self):
        workbook=load_workbook(report.OUTPUT,read_only=True,data_only=True)
        self.assertEqual(workbook.sheetnames,report.SHEETS)
        for s in workbook:self.assertGreater(s.max_row,1)
        workbook.close()
        for name,data in [('01_EXPORT_SIZE',self.a),('02_PROFIT_MANUFACTURING',self.b)]:
            actual=pd.read_excel(report.OUTPUT,sheet_name=name).fillna('')
            expected=report.regression_table(data).fillna('')
            pd.testing.assert_frame_equal(actual,expected,check_dtype=False)
            self.assertFalse(any('Raw' in x or 'Baseline' in x and 'Std' not in x for x in actual.columns[2:]))
        summary=pd.read_excel(report.OUTPUT,sheet_name='03_INTERACTION_SUMMARY')
        self.assertEqual(len(summary),48)
        for column in self.summary:
            if pd.api.types.is_numeric_dtype(self.summary[column]) and not pd.api.types.is_bool_dtype(self.summary[column]):
                np.testing.assert_allclose(summary[column],self.summary[column],rtol=1e-10,atol=1e-10,equal_nan=True)

    def test_constant_manufacturing_cannot_be_accepted(self):
        model=self.b['scenarios']['ALL']['model_results'][1]
        matrix=pd.DataFrame(model['result'].model.exog,columns=model['result'].model.exog_names)
        source,product=report.verify_profit_identifiability(model)
        self.assertEqual(np.linalg.matrix_rank(matrix),len(matrix.columns))
        constant=matrix.copy();constant['manufacturing']=1.;constant[product]=constant[source]
        self.assertLess(np.linalg.matrix_rank(constant),len(constant.columns))

    def test_primary_diagnostics_data_and_shared_configuration_unchanged(self):
        self.assertEqual(self.hashes,{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in self.protected})
        self.assertNotIn('profit_margin_x_manufacturing',engine.INTERACTION_METADATA)


if __name__=='__main__':unittest.main(verbosity=2)
