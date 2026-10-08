"""Compact supplementary OLS interactions; never writes primary/diagnostic reports."""
from __future__ import annotations

from copy import deepcopy
from contextlib import redirect_stdout
import io
import hashlib
import shutil
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import t as student_t, norm

import code_ols_scenarios as engine

OUTPUT = Path("results_ols_interactions.xlsx")
SHEETS = ["00_README", "01_EXPORT_SIZE", "02_PROFIT_MANUFACTURING", "03_INTERACTION_SUMMARY"]
SCENARIOS = ["ALL", "RANK2019", "MANUFACTURING", "RANK2019_MANUFACTURING"]
SAMPLE_LABELS = {"ALL": "ALL", "RANK2019": "RANK2019", "MANUFACTURING": "ALL_MANUFACTURING", "RANK2019_MANUFACTURING": "RANK2019_MANUFACTURING"}
DISPLAY_VARIANTS = ["baseline_std", "winsor_std"]
PROFIT_METADATA = {
    "profit_margin_x_manufacturing": {
        "variables": ["profit_margin", "manufacturing"],
        "display_name": "profit_margin × manufacturing",
        "interpretation": "Difference in the full-sample-standardised profitability slope between manufacturing and other firms.",
        "standardise": False, "scale_with": "profit_margin", "include": True,
    },
}


def specification_config(specification, config=engine.CONFIG):
    result = deepcopy(config)
    result["include_interactions"] = True
    result["analysis_name"] = f"ols_interaction_{specification}"
    if specification == "export_size":
        result["interaction_metadata"] = {k: deepcopy(v) for k, v in engine.INTERACTION_METADATA.items() if v.get("include") and v.get("variables") == ["export_ratio", "ln_sales"]}
        if len(result["interaction_metadata"]) != 1:
            raise ValueError("Exactly one active export_ratio × ln_sales metadata definition is required; shared metadata is not modified.")
    elif specification == "profit_manufacturing":
        result["interaction_metadata"] = deepcopy(PROFIT_METADATA)
        result["base_regressors"] = [*result["base_regressors"], "manufacturing"]
        result["regressor_metadata"] = {
            **result.get("regressor_metadata", {}),
            "manufacturing": {"column_pattern": "manufacturing", "display_name": "Manufacturing", "standardise": False,
                              "variable_type": "dummy", "interpretation": "Manufacturing classification (1), non-manufacturing reference (0); sector controls retained."},
        }
    else:
        raise ValueError(f"Unknown separate specification: {specification}")
    return result


def fit_specification(specification, config=engine.CONFIG):
    lightweight = specification_config(specification, config)
    engine.validate_user_config(lightweight)
    configured = engine.normalise_config(lightweight)
    engine.validate_config(configured)
    models = engine.build_models(configured)
    registry = engine.build_variable_registry(configured, models)
    data = engine.load_input_data(configured)
    engine.validate_input_columns(data, configured, models)
    scenarios = SCENARIOS if specification == "export_size" else ["ALL", "RANK2019"]
    results = {}
    for scenario in scenarios:
        details = engine.run_models_for_scenario(data, scenario, engine.SCENARIOS[scenario], configured, models, registry)
        if details["models_skipped"] or details["models_estimated"] != len(models) * len(engine.get_model_variants(configured)):
            raise ValueError(f"Incomplete supplementary model grid: {specification}/{scenario}; {details['warnings']}")
        results[scenario] = details
    return {"config": configured, "models": models, "scenarios": results}


def paired_reduced_model(report, scenario, model):
    """Refit without the product, keeping its constituent main effects and exact IDs."""
    details = report["scenarios"][scenario]
    c = {**report["config"], "include_interactions": False}
    period = model["period"]
    products = engine.get_interaction_column_names(report["config"])
    spec = {**report["models"][period], "regressors": [x for x in report["models"][period]["regressors"] if x not in products]}
    index = model["result"].model.data.row_labels
    frame = details["filtered_df"].loc[index].copy()
    levels = engine.get_categorical_levels(details["filtered_df"], c)
    variant = next(v for v in engine.get_model_variants(c) if engine.get_model_name(period, v) == model["model"])
    reduced = engine.run_model_variant(period, spec, variant, frame, levels, details["variable_registry"], c)
    if reduced["estimation_sample_ids"] != model["estimation_sample_ids"]:
        raise ValueError("Matched reduced-model observations differ; report not published.")
    np.testing.assert_array_equal(reduced["result"].model.endog, model["result"].model.endog)
    return reduced


def linear_contrast(fit, weights):
    """Conditional slope, SE, p and CI from the FULL estimated covariance matrix."""
    vector = pd.Series(0., index=fit.params.index)
    for column, value in weights.items(): vector[column] = value
    estimate = float(vector @ fit.params)
    variance = float(vector @ fit.cov_params() @ vector)
    if variance < -1e-12:
        raise ValueError("Negative contrast variance; conditional slope not reported.")
    se = float(np.sqrt(max(variance, 0.)))
    statistic = estimate / se if se else (0. if estimate == 0 else np.inf)
    distribution = student_t(df=fit.df_resid) if fit.use_t else norm
    p = float(2 * distribution.sf(abs(statistic)))
    critical = float(distribution.ppf(.975))
    return {"coefficient": estimate, "se": se, "p": p, "ci_low": estimate - critical * se, "ci_high": estimate + critical * se}


def verify_profit_identifiability(model):
    fit = model["result"]
    profit = f"profit_margin_start_{engine.get_regressor_period_for_model(model['period'])}"
    product = engine.build_interaction_column_name("profit_margin_x_manufacturing", model["period"])
    matrix = pd.DataFrame(fit.model.exog, columns=fit.model.exog_names)
    if set(matrix.manufacturing.unique()) != {0., 1.}:
        raise ValueError("Profitability × manufacturing requires both groups; manufacturing-only models are not estimated.")
    rank = np.linalg.matrix_rank(matrix.to_numpy())
    without_terms = matrix.drop(columns=["manufacturing", product])
    if rank != matrix.shape[1] or rank - np.linalg.matrix_rank(without_terms.to_numpy()) != 2:
        raise ValueError("Manufacturing main effect or interaction is not identifiable alongside sector controls.")
    if model["standardised_model"] == "Yes":
        np.testing.assert_allclose(matrix[product], matrix[profit] * matrix.manufacturing, rtol=1e-12, atol=1e-12)
    return profit, product


def summary_table(export, profit):
    rows = []
    for specification, report in [("Export × size", export), ("Profitability × manufacturing", profit)]:
        for scenario in SCENARIOS:
            if scenario not in report["scenarios"]: continue
            models = report["scenarios"][scenario]["model_results"]
            for variant in DISPLAY_VARIANTS:
                for period in report["config"]["periods"]:
                    model = next(m for m in models if m["model"] == f"{period}_{variant}")
                    fit = model["result"]
                    reduced = paired_reduced_model(report, scenario, model)
                    term = engine.get_active_interactions(report["config"])[0]
                    product = engine.build_interaction_column_name(term["name"], period)
                    row = {"interaction_specification": specification, "sample": SAMPLE_LABELS[scenario], "period": period,
                           "model_specification": "Standardised baseline" if variant == "baseline_std" else "Standardised winsorised",
                           "interaction_coefficient": float(fit.params[product]), "interaction_p": float(fit.pvalues[product]),
                           "N": int(fit.nobs), "R_squared": fit.rsquared, "adjusted_R_squared": fit.rsquared_adj,
                           "adjusted_R_squared_without_interaction": reduced["result"].rsquared_adj,
                           "identical_observations": True}
                    if report is profit:
                        source, _ = verify_profit_identifiability(model)
                        nonmanufacturing = linear_contrast(fit, {source: 1.})
                        manufacturing = linear_contrast(fit, {source: 1., product: 1.})
                        row.update({"profitability_non_manufacturing": nonmanufacturing["coefficient"],
                                    "profitability_non_manufacturing_p": nonmanufacturing["p"],
                                    "profitability_manufacturing": manufacturing["coefficient"],
                                    "profitability_manufacturing_p": manufacturing["p"],
                                    "profitability_manufacturing_ci_low": manufacturing["ci_low"],
                                    "profitability_manufacturing_ci_high": manufacturing["ci_high"]})
                    rows.append(row)
    order = ["interaction_specification", "sample", "period", "model_specification", "profitability_non_manufacturing", "profitability_non_manufacturing_p", "profitability_manufacturing", "profitability_manufacturing_p", "interaction_coefficient", "interaction_p", "N", "R_squared", "adjusted_R_squared", "adjusted_R_squared_without_interaction", "profitability_manufacturing_ci_low", "profitability_manufacturing_ci_high", "identical_observations"]
    return pd.DataFrame(rows).reindex(columns=order)


def regression_table(report):
    tables = []
    for scenario in SCENARIOS:
        if scenario not in report["scenarios"]: continue
        details = report["scenarios"][scenario]
        c = report["config"]
        coefficients = details["coefficients_internal_df"].drop(columns="scenario")
        summaries = details["summary_df"].drop(columns="scenario")
        row_order = engine.build_comparison_row_order(c, coefficients, details["variable_registry"])
        table = engine.build_comparison_sheet(coefficients, summaries, c, row_order, DISPLAY_VARIANTS)
        # Same familiar row/column layout; retain dependent-variable identification.
        for label in ["Dependent variable column", "Dependent variable"]:
            values = {}
            for variant in DISPLAY_VARIANTS:
                for period in c["periods"]:
                    meta = engine.period_dependent_metadata(c["growth_mode"], period)
                    values[engine.get_comparison_column_label(f"{period}_{variant}")] = meta["column"] if label.endswith("column") else meta["label"]
            table = pd.concat([pd.DataFrame([{"display_name": label, **values}]), table], ignore_index=True)
        table.insert(0, "sample", SAMPLE_LABELS[scenario])
        tables.append(table)
    return pd.concat(tables, ignore_index=True)


def readme(export, profit):
    c = export["config"]
    rows = [
        ("Purpose", "Supplementary interaction tests; primary additive models remain in results_ols_scenarios.xlsx. No primary or diagnostics workbook is regenerated."),
        ("Samples", "ALL: all eligible complete trajectories. RANK2019: in_rank_2019=1. ALL_MANUFACTURING: manufacturing=1 (shared engine key MANUFACTURING). RANK2019_MANUFACTURING: both filters."),
        ("Export × size", "All four samples; existing export_ratio × ln_sales product and estimation/scaling unchanged. Continuous inputs centred on each estimation sample; product independently standardised in these displayed standardised models."),
        ("Profitability × manufacturing", "ALL and RANK2019 only. Separate model: profit_margin + Manufacturing + profit_margin × manufacturing + existing controls. Export × size is NOT included in this specification. Sector indicators remain."),
        ("Profitability scaling", "Use one full-sample profitability mean/SD for both groups. Product = z(profit_margin) × manufacturing; manufacturing stays 0/1. The product is not divided by its own SD. This new specification's explicit constituent scaling makes beta_profit + beta_interaction the manufacturing slope."),
        ("Interpretation", "Standardised growth change per one full-sample SD of profitability: beta_profit for non-manufacturing, beta_profit + beta_interaction for manufacturing. Difference is beta_interaction; its p-value tests slope equality. Export × size's product coefficient has its established independent product-SD units and is not a group slope difference."),
        ("Scope of slopes", "One full-sample profitability SD can be much larger than a group's own typical variation. Large conditional slopes are not typical within-group changes. These are conditional associations, not causal effects; no predictor clipping or new data-quality exclusions is introduced."),
        ("Inference", f"Retain actual shared covariance estimator: {c['covariance_type']}. No switch to a different robust convention. Manufacturing-slope p/CI uses var(beta_profit)+var(beta_interaction)+2cov(beta_profit,beta_interaction); t or normal reference follows fitted covariance settings."),
        ("Periods", "P1 2019–2020 (1 year); P2 2020–2022 (2 years); P3 2022–2024 (2 years); FULL 2019–2024 (5 years). FULL uses 2019 predictors; P1/P2/P3 use 2019/2020/2022 respectively."),
        ("Dependent variable", f"{c['growth_mode'].title()} annualised log sales growth. Source columns and intervals appear in regression tables. No new outcome definition."),
        ("Controls", "Existing size, profitability, exports, asset turnover, capital ratio, sales per employee, ownership, sector indicators and period-specific lags retained. Production sector reference; FULL has no lag. Profitability specification additionally includes the independent manufacturing main effect."),
        ("Models", f"Baseline before winsorised. Only standardised variants displayed; raw variants remain computed in the engine for verification. Winsorisation: outcome only, {c['winsor_lower']:.0%}/{c['winsor_upper']:.0%} within each estimation sample. Predictors are not clipped."),
        ("Matched comparisons", "Adjusted R² without interaction comes from an actual reduced-model refit on exactly the interaction model's ordered firms and transformed outcome. Manufacturing remains a main effect in the profitability model's comparator. These comparators may therefore differ from the main workbook's additive specification."),
        ("Identification", "Verify full design rank and a two-column rank gain from manufacturing and its product alongside controls/sector indicators. Do not estimate profitability × manufacturing in manufacturing-only samples."),
        ("Presentation", "Coefficient with significance stars above p-value in parentheses. *** p<.01; ** p<.05; * p<.10. Numeric summary fields retain full precision; display four decimals. Empty conditional-profitability cells for export × size are not applicable."),
    ]
    rows += engine.manual_exclusion_readme_rows(c.get("manual_exclusion_audit", []))
    return pd.DataFrame(rows, columns=["item", "description"])


def write_compact_workbook(export, profit, destination=OUTPUT):
    # Complete all comparisons/identification checks before publishing.
    summary = summary_table(export, profit)
    tables = {"00_README": readme(export, profit), "01_EXPORT_SIZE": regression_table(export),
              "02_PROFIT_MANUFACTURING": regression_table(profit), "03_INTERACTION_SUMMARY": summary}
    if list(tables) != SHEETS or any(t.empty for t in tables.values()):
        raise ValueError("The four supplementary sheets must all contain results.")
    destination = Path(destination)
    temporary = destination.with_name(destination.stem + ".writing.xlsx")
    with pd.ExcelWriter(temporary, engine="xlsxwriter") as writer:
        for name, table in tables.items(): table.to_excel(writer, sheet_name=name, index=False)
        workbook = writer.book
        header = workbook.add_format({"font_name":"Arial","font_size":10,"bold":True,"bg_color":"#1F4E78","font_color":"white","text_wrap":True,"valign":"vcenter"})
        body = workbook.add_format({"font_name":"Arial","font_size":10,"text_wrap":True,"valign":"vcenter"})
        decimal = workbook.add_format({"font_name":"Arial","font_size":10,"num_format":"0.0000"})
        p_format = workbook.add_format({"font_name":"Arial","font_size":10,"num_format":'[<0.0001]"<0.0001";0.0000'})
        for name, table in tables.items():
            sheet=writer.sheets[name];sheet.hide_gridlines(2);sheet.set_tab_color('#5B9BD5');sheet.set_zoom(90)
            sheet.freeze_panes(1,2 if name in SHEETS[1:3] else 4 if name==SHEETS[3] else 0)
            sheet.autofilter(0,0,len(table),len(table.columns)-1);sheet.set_row(0,48 if name==SHEETS[3] else 36)
            for i,column in enumerate(table.columns):
                sheet.write(0,i,column,header)
                if name=='00_README':sheet.set_column(i,i,32 if i==0 else 100,body)
                elif name in SHEETS[1:3]:sheet.set_column(i,i,28 if i==0 else 34 if i==1 else 24,body)
                elif pd.api.types.is_numeric_dtype(table[column]) and not pd.api.types.is_bool_dtype(table[column]):sheet.set_column(i,i,21,p_format if column.endswith('_p') else decimal)
                else:sheet.set_column(i,i,10 if column=='period' else 24 if column=='model_specification' else 28,body)
            if name in SHEETS[1:3]:
                for row,label in enumerate(table.display_name,1):sheet.set_row(row,42 if label.startswith('Dependent') else 30 if label not in engine.SUMMARY_ROWS else 18)
            if name=='00_README':
                for row,text in enumerate(table.description,1):sheet.set_row(row,15 * max(2,int(np.ceil(len(text)/95))))
            if name==SHEETS[3]:
                sheet.set_column(table.columns.get_loc('N'),table.columns.get_loc('N'),10,workbook.add_format({'num_format':'0'}))
        writer.book.set_properties({'title':'GRIP supplementary OLS interactions'})
    temporary.replace(destination)
    return tables


def run_interactions(config=engine.CONFIG, export_report=None, destination=OUTPUT):
    # Independent entry point deliberately does not call the dual primary writer.
    protected = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in Path('.').rglob('*') if p.is_file() and p.suffix in {'.xlsx','.parquet'} and p.resolve()!=Path(destination).resolve()}
    with redirect_stdout(io.StringIO()):
        export = export_report or fit_specification('export_size',config)
        profit = fit_specification('profit_manufacturing',config)
    # Keep the prior full report, including reader annotations, as a historical
    # record before replacing its active presentation with the four-sheet view.
    destination=Path(destination)
    if destination.exists():
        from openpyxl import load_workbook
        old=load_workbook(destination,read_only=True)
        was_full='Coefficients_Long' in old.sheetnames
        old.close()
        if was_full:
            archive=destination.parent/'archive'/f'results_ols_interactions_{date.today().isoformat()}.xlsx'
            sequence=2
            while archive.exists():
                archive=destination.parent/'archive'/f'results_ols_interactions_snapshot_{sequence}_{date.today().isoformat()}.xlsx'
                sequence+=1
            archive.parent.mkdir(exist_ok=True)
            shutil.copy2(destination,archive)
    tables=write_compact_workbook(export,profit,destination)
    if any(hashlib.sha256(p.read_bytes()).hexdigest()!=h for p,h in protected.items()):
        raise RuntimeError('Protected workbook/data changed during the run; investigate concurrent external changes. This runner only writes its interaction report.')
    print('Compact interaction report: 4 sheets; 32 export × size and 16 profitability × manufacturing displayed models.')
    print('Manufacturing-only samples retain export × size only; both separate specifications and raw engine variants preserved.')
    print(tables['03_INTERACTION_SUMMARY'].query("interaction_specification == 'Profitability × manufacturing'")[["sample","period","model_specification","interaction_coefficient","interaction_p","N"]].to_string(index=False))
    return {'export':export,'profit':profit,'tables':tables}


if __name__=='__main__':
    run_interactions()
