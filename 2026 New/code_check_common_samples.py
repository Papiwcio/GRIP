"""Research-grade checks of central eligibility, estimator identities and reports."""
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import code_common_samples as common


def validate_context(context):
    assert not context.source.nip.duplicated().any()
    for population,ids in context.memberships.items():
        assert len(ids)==len(set(ids))
        assert ids <= context.scopes[population]
        assert not ids & context.manual
        for (p,spec),eligible in context.eligible.items():
            if p==population: assert ids <= eligible
        blocked=context.details.loc[context.details.eligibility_impact & context.details.populations.str.split('; ').apply(lambda p:population in p),'nip']
        assert not ids & set(blocked)
    assert context.memberships['MANUFACTURING'] <= context.memberships['ALL']
    assert context.memberships['RANK2019_MANUFACTURING'] <= context.memberships['RANK2019']
    assert not context.details.duplicated(['nip','year','variable','rule']).any()
    assert not ((context.details.year==2018)&context.details.variable.ne('sales')&context.details.eligibility_impact).any()


def validate_workbooks(directory,context):
    directory=Path(directory)
    metadata=['population_id','sample_id','specification_id','model_id','unique_companies','sample_sha256']
    primary=pd.read_excel(directory/'results_ols_scenarios.xlsx',sheet_name='Model_Summary_Long')
    quantile=pd.read_excel(directory/'results_quantile.xlsx',sheet_name='Model_Summary_Long')
    supplement=pd.read_excel(directory/'results_ols_interactions.xlsx',sheet_name='03_INTERACTION_SUMMARY')
    for table,count in [(primary,64),(quantile,48),(supplement,48)]:
        assert len(table)==count and set(metadata)<=set(table)
        for population,rows in table.groupby('population_id'):
            selected=context.memberships[population]
            assert set(rows.sample_sha256)=={common.fingerprint(selected)}
            assert set(rows.unique_companies)=={len(selected)}
    correlation=pd.read_excel(directory/'results_diagnostics_trajectories.xlsx',sheet_name='90_CORR_LONG_ALL')
    for population,rows in correlation.groupby('scenario'):
        assert set(rows.N)=={len(context.memberships[population])}
    wb=load_workbook(directory/'results_data_quality_and_samples.xlsx',read_only=False,data_only=True)
    from code_report_quality_samples import SHEETS
    assert wb.sheetnames==SHEETS
    wb.close()
    excluded=pd.read_excel(directory/'results_data_quality_and_samples.xlsx',sheet_name='03_EXCLUDED_COMPANIES',header=3,dtype={'nip':str})
    detail=pd.read_excel(directory/'results_data_quality_and_samples.xlsx',sheet_name='04_EXCLUSION_DETAILS',header=3,dtype={'nip':str})
    member=pd.read_excel(directory/'results_data_quality_and_samples.xlsx',sheet_name='05_SAMPLE_MEMBERSHIP',header=3,dtype={'nip':str})
    assert not excluded.nip.duplicated().any()
    assert not detail.duplicated(['nip','year','variable','rule']).any()
    assert not member.nip.duplicated().any() and len(member)==len(context.source)
    for population,ids in context.memberships.items():
        assert set(member.loc[member[population].eq('INCLUDED'),'nip'])==ids
    assert 'Dropped_Rows_Long' not in pd.ExcelFile(directory/'results_ols_scenarios.xlsx').sheet_names
    print('PASS: common eligibility, nesting, 160 saved model metadata rows, correlations and consolidated registers.',flush=True)


def regression_math_equivalence(context):
    import importlib.util
    import code_ols_scenarios as engine
    import code_ols_interactions as supplement
    reference=Path('/tmp/grip_pre_stage2_ols_engine.py')
    if not reference.exists():
        import subprocess
        content=subprocess.check_output(['git','show','674476e:2026 New/code_ols_scenarios.py'],cwd=common.ROOT)
        reference.write_bytes(content)
    spec=importlib.util.spec_from_file_location('grip_legacy_reference',reference)
    legacy=importlib.util.module_from_spec(spec);spec.loader.exec_module(legacy)
    common.activate(context)
    for family,config in [('additive',{**engine.CONFIG,'include_interactions':False}),('export',supplement.specification_config('export_size')),('profit',supplement.specification_config('profit_manufacturing'))]:
        config=engine.normalise_config(config)
        models=engine.build_models(config)
        frame=common.select(context.source,'RANK2019')
        frame=engine.add_interaction_columns(frame,config,models)
        levels=engine.get_categorical_levels(frame,config)
        registry=engine.build_variable_registry(config,models,levels)
        for period,model in models.items():
            for variant in engine.get_model_variants(config):
                new=engine.run_model_variant(period,model,variant,frame,levels,registry,config)
                old=legacy.run_model_variant(period,model,variant,frame,levels,registry,config)
                assert set(new['estimation_sample_ids'])==context.memberships['RANK2019']
                np.testing.assert_array_equal(new['result'].model.endog,old['result'].model.endog)
                np.testing.assert_allclose(new['result'].model.exog,old['result'].model.exog,rtol=0,atol=0)
                for field in ['params','bse','pvalues']:
                    np.testing.assert_allclose(getattr(new['result'],field),getattr(old['result'],field),rtol=1e-12,atol=1e-12)
    print('PASS: 48 matched-sample fits reproduce the preserved engine exactly (outcomes, design matrices, coefficients, SEs and p-values).',flush=True)


if __name__=='__main__':
    context=common.current();validate_context(context);validate_workbooks(common.ROOT,context);regression_math_equivalence(context)
